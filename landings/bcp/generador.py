#!/usr/bin/env python3
"""Generador de la landing BCP de Juntoz a partir de Google Sheets.

Lee las pestañas del Sheet (o un .xlsx descargado de él), valida el contenido,
lo inyecta en plantilla.html y escribe:
  - salida/bcp-landing.preview.html  vista previa local (siempre, aunque haya errores)
  - publicado/bcp-landing.html       HTML final para pegar en Magento (solo si todo es válido)
  - respaldos/                       copia de la versión publicada anterior + diff

Uso rápido (ver README.md):
  python3 generador.py plantilla                       # crea la plantilla .xlsx para el Sheet
  python3 generador.py previsualizar --fuente xlsx --archivo BCP.xlsx
  python3 generador.py previsualizar                   # lee Google Sheets (BCP_SHEET_ID + GOOGLE_APPLICATION_CREDENTIALS)
  python3 generador.py publicar --verificar-urls

El HTML final no lleva JavaScript, credenciales ni atributos data-field.
"""
from __future__ import annotations

import argparse
import datetime as dt
import difflib
import os
import re
import shutil
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, StrictUndefined
from markupsafe import Markup, escape

BASE = Path(__file__).resolve().parent
PLANTILLA = "plantilla.html"
DIR_SALIDA = BASE / "salida"
DIR_PUBLICADO = BASE / "publicado"
DIR_RESPALDOS = BASE / "respaldos"
NOMBRE_FINAL = "bcp-landing.html"

PESTANAS = ["Config", "Textos", "Slides", "Productos", "Beneficios", "Pasos", "FAQ", "Legales"]
COLUMNAS = {
    "Config": ["clave", "valor", "ayuda"],
    "Textos": ["clave", "valor", "ayuda"],
    "Slides": ["orden", "visible", "titulo", "descripcion", "imagen", "imagen_alt",
               "cta_texto", "cta_enlace", "cta2_texto", "cta2_enlace"],
    "Productos": ["orden", "visible", "marca", "nombre", "imagen", "imagen_alt", "enlace",
                  "beneficio", "etiqueta"],
    "Beneficios": ["orden", "visible", "icono", "titulo", "texto"],
    "Pasos": ["orden", "visible", "titulo", "texto"],
    "FAQ": ["orden", "visible", "pregunta", "respuesta"],
    "Legales": ["orden", "texto"],
}

# Anclas internas de la plantilla que se aceptan como enlace de un botón.
ANCLAS = {"#bcp-landing", "#bcp-productos", "#bcp-como-usar"}

# (clave, obligatorio, tipo, valor por defecto en la plantilla, ayuda)
# Los datos comerciales del banco van vacíos a propósito: no se inventan.
CONFIG = [
    ("banco", True, "texto", "BCP", "Nombre del banco tal como se muestra."),
    ("campana", True, "texto", "", "Nombre interno de la campaña (no se muestra). Ej.: BCP octubre 2026."),
    ("codigo", True, "codigo", "", "Código del cupón, sin espacios. Ej.: BCP200."),
    ("porcentaje", True, "porcentaje", "", "Solo el número. Ej.: 15"),
    ("tope", True, "monto", "", "Descuento máximo en soles por transacción, solo el número. Ej.: 200"),
    ("minimo", True, "monto", "", "Compra mínima en soles, solo el número. Ej.: 399"),
    ("fecha_inicio", True, "fecha", "", "AAAA-MM-DD o DD/MM/AAAA."),
    ("fecha_fin", True, "fecha", "", "AAAA-MM-DD o DD/MM/AAAA. Si ya pasó, no se publica."),
    ("vigencia_texto", False, "texto", "", "Opcional. Si se deja vacío se arma desde las fechas: «del 21 al 28 de octubre de 2026»."),
    ("tarjetas", True, "texto", "", "Tarjetas BCP que aplican, tal como figuran en los legales."),
    ("legales_aprobados", True, "sino", "NO", "SI solo cuando BCP y Juntoz aprobaron el texto de la pestaña Legales."),
    ("legales_aprobado_por", False, "texto", "", "Quién aprobó los legales (obligatorio si legales_aprobados = SI)."),
    ("legales_fecha_aprobacion", False, "fecha", "", "Fecha de aprobación (obligatoria si legales_aprobados = SI)."),
    ("dominios_enlaces", True, "dominios", "juntoz.com", "Dominios permitidos en los enlaces, separados por coma."),
    ("dominios_imagenes", True, "dominios",
     "juntoz.com, juntozstgsrvproduction.blob.core.windows.net, mcprod.efe.com.pe",
     "Dominios permitidos para imágenes, separados por coma."),
    ("color_principal", True, "color", "#0030B2", "Azul BCP: botones, títulos destacados."),
    ("color_principal_hover", True, "color", "#00227F", "Azul al pasar el mouse por los botones."),
    ("color_acento", True, "color", "#FF7800", "Naranja BCP: detalles y foco del teclado."),
    ("color_texto", True, "color", "#171717", "Texto principal."),
    ("color_texto_suave", True, "color", "#5B6272", "Texto secundario."),
    ("color_fondo_secundario", True, "color", "#F5F7FB", "Fondo del hero y de productos."),
    ("color_borde", True, "color", "#E3E8F2", "Bordes de tarjetas."),
]

# (clave, obligatorio, valor por defecto = copy actual de la landing, ayuda)
TEXTOS = [
    ("hero_insignia_marca", True, "{banco}", "Pastilla azul del hero."),
    ("hero_insignia_texto", True, "Beneficio exclusivo en Juntoz", "Texto junto a la pastilla."),
    ("hero_dato1_etiqueta", True, "Descuento", "Dato 1 del hero (etiqueta)."),
    ("hero_dato1_valor", True, "{porcentaje}%", "Dato 1 del hero (valor)."),
    ("hero_dato2_etiqueta", True, "Tope", "Dato 2 del hero (etiqueta)."),
    ("hero_dato2_valor", True, "S/{tope}", "Dato 2 del hero (valor)."),
    ("hero_dato3_etiqueta", True, "Desde", "Dato 3 del hero (etiqueta)."),
    ("hero_dato3_valor", True, "S/{minimo}", "Dato 3 del hero (valor)."),
    ("hero_vigencia", True, "Vigencia: {vigencia}", "Línea bajo los botones del hero."),
    ("ticket_etiqueta", True, "Tu código", "Cupón flotante del hero (si el slide no tiene imagen)."),
    ("ticket_texto", True, "Hasta S/{tope} de dscto.", "Texto bajo el código en el cupón flotante."),
    ("beneficios_etiqueta", True, "Tu beneficio", "Antetítulo de beneficios."),
    ("beneficios_titulo", True, "Así funciona el descuento {banco}", "H2 de beneficios."),
    ("productos_etiqueta", True, "Selección exclusiva", "Antetítulo de productos."),
    ("productos_titulo", True, "Productos para aprovechar tu cupón", "H2 de productos."),
    ("productos_bajada", False, "Aplica {codigo} en compras desde S/{minimo} y llévate hasta {porcentaje}% de descuento.", "Párrafo bajo el H2 de productos."),
    ("producto_beneficio_etiqueta", True, "Con {banco}:", "Texto antes del beneficio de cada producto."),
    ("producto_beneficio_defecto", True, "{porcentaje}% hasta S/{tope}", "Beneficio de un producto cuando su columna «beneficio» está vacía."),
    ("producto_cta", True, "Comprar", "Botón de cada producto."),
    ("pasos_etiqueta", True, "Compra fácil", "Antetítulo de pasos."),
    ("pasos_titulo", True, "Usa tu cupón en {n_pasos} pasos", "H2 de pasos. {n_pasos} = pasos visibles."),
    ("codigo_titulo", True, "Tu código {banco}", "H2 de la franja azul."),
    ("codigo_bajada", False, "Aplícalo en compras desde S/{minimo} y paga con tu tarjeta {banco}.", "Párrafo de la franja azul."),
    ("codigo_ayuda", False, "Selecciona el código para copiarlo", "Ayuda bajo el código."),
    ("codigo_cta_texto", False, "Comprar ahora", "Botón de la franja azul (vacío = sin botón)."),
    ("codigo_cta_enlace", False, "#bcp-productos", "Enlace del botón: #bcp-productos o una URL de Juntoz."),
    ("faq_etiqueta", True, "Ayuda", "Antetítulo de preguntas frecuentes."),
    ("faq_titulo", True, "Preguntas frecuentes", "H2 de preguntas frecuentes."),
    ("legales_titulo", True, "{banco} · Términos y condiciones", "Título de los legales."),
]

TOKENS_AYUDA = "{banco} {codigo} {porcentaje} {tope} {minimo} {vigencia} {tarjetas} {n_pasos}"

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "setiembre", "octubre", "noviembre", "diciembre"]

RE_MARCADOR = re.compile(r"\[(?i:XX|C[ÓO]DIGO|MONTO_\w+|FECHAS|TARJETAS)\]|(?i:lorem ipsum)|\bTODO\b|\bPENDIENTE\b")
RE_TOKEN = re.compile(r"\{([a-z_]+)\}")
RE_NEGRITA = re.compile(r"\*\*(.+?)\*\*")
RE_COLOR = re.compile(r"^#[0-9A-Fa-f]{6}$")
RE_CODIGO = re.compile(r"^[A-Za-z0-9_-]{3,30}$")
# Patrones de secretos que nunca deben terminar en el HTML.
RE_SECRETO = re.compile(r"private_key|BEGIN [A-Z ]*PRIVATE KEY|AIza[0-9A-Za-z_-]{20,}|ya29\.[0-9A-Za-z_-]+|client_secret", re.I)
FALTA_INI, FALTA_FIN = "⟦FALTA: ", "⟧"
IMG_VACIA = ("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 400 400'%3E"
             "%3Crect width='400' height='400' fill='%23F5F7FB'/%3E%3C/svg%3E")

# Datos de ejemplo de la plantilla. Productos y URLs reales de Juntoz (landing BBVA vigente);
# el beneficio de cada producto sale de la configuración de la campaña.
_BLOB = "https://juntozstgsrvproduction.blob.core.windows.net/default-blob-images/"
EJEMPLO = {
    "Slides": [
        [1, "SI", "Ahorra hasta **{porcentaje}%** pagando con tus tarjetas {banco}",
         "Usa el código **{codigo}** y obtén hasta S/{tope} de descuento en compras desde S/{minimo}.",
         "", "", "Ver productos", "#bcp-productos", "¿Cómo uso mi cupón?", "#bcp-como-usar"],
    ],
    "Productos": [
        [1, "SI", "Apple", "PRE-VENTA IPHONE 18 PRO SILVER 512GB", _BLOB + "orig_18prosilver_577686.jpg", "",
         "https://juntoz.com/p/nuevo-iphone-18-pro-silver-512gb-esim-35505903iphn18512slvtt0909202608273082283", "", "Preventa"],
        [2, "SI", "Samsung", "SAMSUNG S26 ULTRA 12GB 512GB SKY BLUE", _BLOB + "orig_161_960625.jpg", "",
         "https://juntoz.com/p/samsung-s26-ultra-12gb-512gb-sky-blue-35505903samsung26ultraskyblue5122904202607335392365", "", ""],
        [3, "SI", "LG", "Lavadora LG WT13OBVTB 13Kg Carga Superior – Negro Onyx", _BLOB + "orig_7deab16f706bb7b5e538d306beccdf92_566828.jpg", "",
         "https://juntoz.com/p/lavadora-lg-wt13obvtb-13kg-carga-superior-negro-onyx-35415892lavadoralgwt13obvtb1505202608031327663", "", ""],
        [4, "SI", "HP", 'Laptop HP 255R G10 15.6" AMD Ryzen 5 16GB 512GB', _BLOB + "orig_cz1p1lt1_118442.jpg", "",
         "https://juntoz.com/p/laptop-hp-255r-g10-156-amd-ryzen-5-16gb-512gb-cz1p1ltabm-17703601cz1p1lt0709202609324730151", "", ""],
        [5, "SI", "TCL", 'Televisor Tcl 75" Qled 4k 75p7k Google Tv', _BLOB + "orig_captura-de-pantalla-2026-01-24-183103_193059.jpg", "",
         "https://juntoz.com/p/televisor-tcl-75-qled-4k-75p7k-google-tv-34685816lnmwshjklr2401202611530207106", "", ""],
        [6, "SI", "Epson", "Impresora EPSON Multifunción 3 en 1", _BLOB + "orig_epson-l4360_775599.jpg", "",
         "https://juntoz.com/p/impresora-epson-multifuncion-3-en-1-ecotank-l4360-duplexwi-fi-22324310066763108202509392788312", "", ""],
        [7, "SI", "Sony", "Consola PlayStation 5 Sony Slim con Lectora de Discos", _BLOB + "orig_cfi-2115a_788457.jpg", "",
         "https://juntoz.com/p/consola-playstation-5-sony-slim-con-lectora-de-discos-17703601cfi-2115a1908202609322283889", "", ""],
        [8, "SI", "Hugo Boss", "Hugo Boss Bottled Eau de Toilette 100 ml", _BLOB + "orig_82464578_624025.jpg", "",
         "https://juntoz.com/p/boss-frag-bottled-hombre-edt-100-ml-18603727824645782803202309173378611", "", ""],
    ],
    "Beneficios": [
        [1, "SI", "%", "{porcentaje}% de dscto.", "Hasta S/{tope} por compra."],
        [2, "SI", "S/", "Desde S/{minimo}", "Monto mínimo en el carrito."],
        [3, "SI", "▭", "Tarjetas {banco}", "{tarjetas}."],
        [4, "SI", "◷", "Por tiempo limitado", "Válido {vigencia}."],
    ],
    "Pasos": [
        [1, "SI", "Elige tus productos", "Agrega al carrito desde S/{minimo}."],
        [2, "SI", "Ingresa el código", "Escribe **{codigo}** en el campo de cupones del checkout."],
        [3, "SI", "Verifica tu descuento", "Revisa el {porcentaje}% aplicado antes de pagar."],
        [4, "SI", "Paga con {banco}", "Finaliza con tu tarjeta {banco} participante."],
    ],
    "FAQ": [
        [1, "SI", "¿Cuál es el código de descuento?", "El código es {codigo}. Ingrésalo en el campo de cupones del checkout."],
        [2, "SI", "¿Cuál es el descuento y el monto mínimo?", "{porcentaje}% de descuento hasta S/{tope} por compra, desde S/{minimo}."],
        [3, "SI", "¿Qué tarjetas aplican?", "{tarjetas}."],
        [4, "SI", "¿Cuál es la vigencia?", "La promoción es válida {vigencia}."],
    ],
    "Legales": [[1, ""]],
}


# --------------------------------------------------------------------------- utilidades

def normalizar_clave(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", s.strip().lower()).strip("_")


def celda(v) -> str:
    """Convierte una celda (xlsx o API) a texto, sin tocar la redacción interna."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "SI" if v else "NO"
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    return str(v).replace("\r\n", "\n").replace("\r", "\n").strip()


def a_registros(filas: list[list]) -> list[dict]:
    """Primera fila = encabezados. Devuelve dicts con la fila real del Sheet en '_fila'."""
    if not filas:
        return []
    encabezados = [normalizar_clave(celda(h)) for h in filas[0]]
    salida = []
    for i, fila in enumerate(filas[1:], start=2):
        valores = [celda(v) for v in fila]
        if not any(valores):
            continue
        reg = {h: (valores[j] if j < len(valores) else "") for j, h in enumerate(encabezados) if h}
        reg["_fila"] = i
        salida.append(reg)
    return salida


def plano(s: str) -> str:
    return RE_NEGRITA.sub(r"\1", s)


def rico(s: str, destacado: str = "strong") -> Markup:
    """Escapa HTML; **texto** -> <strong> (o <span> en el H1); saltos de línea -> <br>."""
    t = str(escape(s))
    t = RE_NEGRITA.sub(lambda m: f"<{destacado}>{m.group(1)}</{destacado}>", t)
    return Markup(t.replace("\n", "<br>"))


def literal(s: str) -> Markup:
    """Para legales: solo escapa y respeta saltos de línea. No cambia la redacción."""
    return Markup(str(escape(s)).replace("\n", "<br>"))


def parsear_fecha(s: str):
    s = s.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parsear_numero(s: str):
    t = s.replace("S/", "").replace("s/", "").replace("%", "").replace(",", "").strip()
    try:
        n = float(t)
    except ValueError:
        return None
    return n


def fmt_numero(n: float) -> str:
    return str(int(n)) if float(n).is_integer() else f"{n:.2f}"


def texto_vigencia(ini: dt.date, fin: dt.date) -> str:
    if ini.year != fin.year:
        return f"del {ini.day} de {MESES[ini.month - 1]} de {ini.year} al {fin.day} de {MESES[fin.month - 1]} de {fin.year}"
    if ini.month != fin.month:
        return f"del {ini.day} de {MESES[ini.month - 1]} al {fin.day} de {MESES[fin.month - 1]} de {fin.year}"
    return f"del {ini.day} al {fin.day} de {MESES[fin.month - 1]} de {fin.year}"


def es_visible(s: str):
    v = normalizar_clave(s)
    if v in ("si", "s", "true", "verdadero", "1", "x", "yes"):
        return True
    if v in ("", "no", "n", "false", "falso", "0"):
        return False
    return None


# --------------------------------------------------------------------------- lectura

def leer_xlsx(ruta: Path) -> dict[str, list[list]]:
    from openpyxl import load_workbook
    wb = load_workbook(ruta, data_only=True, read_only=True)
    datos = {}
    for p in PESTANAS:
        if p in wb.sheetnames:
            datos[p] = [list(r) for r in wb[p].iter_rows(values_only=True)]
    return datos


def leer_sheets(sheet_id: str, credenciales: str) -> dict[str, list[list]]:
    """Lee las pestañas con la API de Google Sheets (solo lectura) y una cuenta de servicio."""
    try:
        from google.auth.transport.requests import AuthorizedSession
        from google.oauth2 import service_account
    except ImportError:
        sys.exit("Falta la librería google-auth. Instala: pip install -r requirements.txt")
    ruta = Path(credenciales).expanduser().resolve()
    if not ruta.is_file():
        sys.exit(f"No encuentro el archivo de credenciales: {ruta}")
    repo = BASE.parent.parent.resolve()
    if repo in ruta.parents:
        sys.exit("El archivo de credenciales está dentro del repo. Muévelo fuera (por ejemplo a ~/.config/juntoz/) para no subirlo por error.")
    creds = service_account.Credentials.from_service_account_file(
        str(ruta), scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
    sesion = AuthorizedSession(creds)
    params = [("ranges", f"'{p}'") for p in PESTANAS] + [("valueRenderOption", "FORMATTED_VALUE")]
    r = sesion.get(f"https://sheets.googleapis.com/v4/spreadsheets/{sheet_id}/values:batchGet",
                   params=params, timeout=30)
    if r.status_code in (403, 404):
        sys.exit(f"Google respondió {r.status_code}: comparte el Sheet (solo lectura) con el correo de la cuenta de servicio "
                 f"({creds.service_account_email}) y revisa el ID.")
    if r.status_code == 400:
        sys.exit(f"Google respondió 400. Revisa que existan las pestañas: {', '.join(PESTANAS)}.\n{r.text[:300]}")
    r.raise_for_status()
    rangos = r.json().get("valueRanges", [])
    return {p: vr.get("values", []) for p, vr in zip(PESTANAS, rangos)}


# --------------------------------------------------------------------------- construcción y validación

@dataclass
class Informe:
    errores: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    def error(self, msg):
        self.errores.append(msg)

    def aviso(self, msg):
        self.avisos.append(msg)


def dominio_ok(url: str, dominios: list[str]) -> bool:
    host = (urllib.parse.urlparse(url).hostname or "").lower()
    return any(host == d or host.endswith("." + d) for d in dominios)


def revisar_url(url: str, dominios: list[str], donde: str, inf: Informe, ancla: bool = False) -> bool:
    if ancla and url in ANCLAS:
        return True
    if url.startswith("#"):
        inf.error(f"{donde}: el ancla «{url}» no existe en la landing. Usa {', '.join(sorted(ANCLAS))}.")
        return False
    p = urllib.parse.urlparse(url)
    if p.scheme != "https" or not p.hostname or " " in url:
        inf.error(f"{donde}: «{url}» no es una URL https válida.")
        return False
    if not dominio_ok(url, dominios):
        inf.error(f"{donde}: el dominio de «{url}» no está en la lista permitida ({', '.join(dominios)}).")
        return False
    return True


def construir(datos: dict[str, list[list]], hoy: dt.date | None = None):
    """Devuelve (contexto para la plantilla, informe, urls a verificar en red)."""
    hoy = hoy or dt.date.today()
    inf = Informe()
    urls: dict[str, tuple[bool, str]] = {}  # url -> (es_imagen, dónde)

    for p in PESTANAS:
        if p not in datos:
            inf.error(f"Falta la pestaña «{p}» en el Sheet.")
    regs = {p: a_registros(datos.get(p, [])) for p in PESTANAS}

    for p, filas in regs.items():
        for r in filas:
            for k, v in r.items():
                if k != "_fila" and k != "ayuda" and isinstance(v, str) and RE_MARCADOR.search(v):
                    inf.error(f"{p} fila {r['_fila']}, {k}: tiene un marcador sin completar («{RE_MARCADOR.search(v).group(0)}»).")

    def falta(donde: str) -> str:
        inf.error(f"{donde}: está vacío y es obligatorio.")
        return f"{FALTA_INI}{donde}{FALTA_FIN}"

    # ---- Config
    cfg_raw = {normalizar_clave(r.get("clave", "")): r.get("valor", "") for r in regs["Config"]}
    c: dict[str, str] = {}
    fechas = {}
    for clave, oblig, tipo, defecto, _ in CONFIG:
        v = cfg_raw.get(clave, "")
        donde = f"Config, {clave}"
        if tipo == "color":
            if not v:
                inf.aviso(f"{donde}: vacío, se usa el color por defecto {defecto}.")
                v = defecto
            elif not RE_COLOR.match(v):
                inf.error(f"{donde}: «{v}» no es un color #RRGGBB.")
                v = defecto
            c[clave] = v
            continue
        if not v:
            c[clave] = falta(donde) if oblig else ""
            continue
        if tipo == "codigo" and not RE_CODIGO.match(v):
            inf.error(f"{donde}: «{v}» no es un código válido (3 a 30 letras, números, - o _; sin espacios).")
        elif tipo in ("porcentaje", "monto"):
            n = parsear_numero(v)
            if n is None or n <= 0 or (tipo == "porcentaje" and n > 100):
                inf.error(f"{donde}: «{v}» no es un número válido.")
            else:
                v = fmt_numero(n)
        elif tipo == "fecha":
            f = parsear_fecha(v)
            if not f:
                inf.error(f"{donde}: «{v}» no es una fecha válida (AAAA-MM-DD o DD/MM/AAAA).")
            fechas[clave] = f
        elif tipo == "sino" and es_visible(v) is None:
            inf.error(f"{donde}: debe ser SI o NO.")
        c[clave] = v

    ini, fin = fechas.get("fecha_inicio"), fechas.get("fecha_fin")
    if ini and fin:
        if fin < ini:
            inf.error("Config: fecha_fin es anterior a fecha_inicio.")
        if fin < hoy:
            inf.error(f"Config: la promoción venció el {fin.isoformat()}. No se publica una campaña vencida.")
        if ini > hoy:
            inf.aviso(f"Config: la promoción empieza el {ini.isoformat()}; publícala solo cuando corresponda.")
    vigencia = c["vigencia_texto"] or (texto_vigencia(ini, fin) if ini and fin else f"{FALTA_INI}Config, fechas{FALTA_FIN}")

    if es_visible(c["legales_aprobados"]) is not True:
        inf.error("Config, legales_aprobados: no está en SI. No se publica sin legales aprobados.")
    else:
        if not c["legales_aprobado_por"]:
            inf.error("Config, legales_aprobado_por: obligatorio cuando los legales están aprobados.")
        if not c["legales_fecha_aprobacion"]:
            inf.error("Config, legales_fecha_aprobacion: obligatoria cuando los legales están aprobados.")

    dom_enlaces = [d.strip().lower() for d in c["dominios_enlaces"].split(",") if d.strip() and FALTA_INI not in d]
    dom_imagenes = [d.strip().lower() for d in c["dominios_imagenes"].split(",") if d.strip() and FALTA_INI not in d]

    # ---- listas con orden y visibilidad
    def visibles(pestana: str, con_visible: bool = True) -> list[dict]:
        filas = []
        for r in regs[pestana]:
            donde = f"{pestana} fila {r['_fila']}"
            if con_visible:
                vis = es_visible(r.get("visible", ""))
                if vis is None:
                    inf.error(f"{donde}, visible: «{r.get('visible')}» debe ser SI o NO.")
                    continue
                if not vis:
                    if not r.get("visible"):
                        inf.aviso(f"{donde}: «visible» vacío, la fila no se muestra.")
                    continue
            n = parsear_numero(r.get("orden", ""))
            if n is None or not float(n).is_integer():
                inf.error(f"{donde}, orden: «{r.get('orden', '')}» debe ser un número entero.")
                continue
            r["_orden"] = int(n)
            r["ref"] = donde
            filas.append(r)
        ordenes = [r["_orden"] for r in filas]
        repetidos = sorted({o for o in ordenes if ordenes.count(o) > 1})
        if repetidos:
            inf.aviso(f"{pestana}: orden repetido {repetidos}; se respeta el orden de las filas.")
        return sorted(filas, key=lambda r: (r["_orden"], r["_fila"]))

    slides_r, productos_r = visibles("Slides"), visibles("Productos")
    beneficios_r, pasos_r, faq_r = visibles("Beneficios"), visibles("Pasos"), visibles("FAQ")
    legales_r = visibles("Legales", con_visible=False)

    tokens = {
        "banco": c["banco"], "codigo": c["codigo"], "porcentaje": c["porcentaje"], "tope": c["tope"],
        "minimo": c["minimo"], "vigencia": vigencia, "tarjetas": c["tarjetas"], "n_pasos": str(len(pasos_r)),
    }

    def expandir(s: str, donde: str) -> str:
        def rep(m):
            if m.group(1) not in tokens:
                inf.error(f"{donde}: variable desconocida «{{{m.group(1)}}}». Variables válidas: {TOKENS_AYUDA}.")
                return m.group(0)
            return tokens[m.group(1)]
        return RE_TOKEN.sub(rep, s)

    def campo(r: dict, col: str, oblig: bool) -> str:
        v = r.get(col, "")
        if not v:
            return falta(f"{r['ref']}, {col}") if oblig else ""
        return expandir(v, f"{r['ref']}, {col}")

    # ---- Textos
    t_raw = {normalizar_clave(r.get("clave", "")): r.get("valor", "") for r in regs["Textos"]}
    t: dict[str, Markup | str] = {}
    for clave, oblig, _, _ in TEXTOS:
        v = t_raw.get(clave, "")
        donde = f"Textos, {clave}"
        t[clave] = expandir(v, donde) if v else (falta(donde) if oblig else "")
    if t["codigo_cta_texto"]:
        if not t["codigo_cta_enlace"]:
            t["codigo_cta_enlace"] = "#"
            falta("Textos, codigo_cta_enlace")
        else:
            revisar_url(t["codigo_cta_enlace"], dom_enlaces, "Textos, codigo_cta_enlace", inf, ancla=True)
    t["producto_cta_plano"] = plano(t["producto_cta"])
    enlace_cta = t["codigo_cta_enlace"]
    t = {k: (rico(v) if k not in ("codigo_cta_enlace", "producto_cta_plano") else v) for k, v in t.items()}
    t["codigo_cta_enlace"] = enlace_cta
    datos_hero = [{"clave": f"hero_dato{i}", "etiqueta": t[f"hero_dato{i}_etiqueta"], "valor": t[f"hero_dato{i}_valor"]}
                  for i in (1, 2, 3)]

    def imagen(r: dict, col: str, oblig: bool) -> str:
        url = r.get(col, "")
        donde = f"{r['ref']}, {col}"
        if not url:
            if oblig:
                falta(donde)
                return IMG_VACIA
            return ""
        if revisar_url(url, dom_imagenes, donde, inf):
            urls.setdefault(url, (True, donde))
            return url
        return IMG_VACIA

    def enlace(r: dict, col: str) -> str:
        url = r.get(col, "")
        donde = f"{r['ref']}, {col}"
        if not url:
            falta(donde)
            return "#"
        if revisar_url(url, dom_enlaces, donde, inf, ancla=True):
            if not url.startswith("#"):
                urls.setdefault(url, (False, donde))
            return url
        return "#"

    # ---- Slides
    slides = []
    for r in slides_r:
        titulo = campo(r, "titulo", True)
        s = {"ref": r["ref"], "titulo": rico(titulo, "span"), "descripcion": rico(campo(r, "descripcion", False))}
        s["imagen"] = imagen(r, "imagen", False)
        s["imagen_alt"] = r.get("imagen_alt") or plano(titulo)
        if s["imagen"] and not r.get("imagen_alt"):
            inf.aviso(f"{r['ref']}, imagen_alt: vacío, se usa el título como texto alternativo.")
        for n in ("cta", "cta2"):
            s[f"{n}_texto"] = rico(campo(r, f"{n}_texto", False))
            s[f"{n}_enlace"] = enlace(r, f"{n}_enlace") if s[f"{n}_texto"] else ""
        if len(plano(titulo)) > 90:
            inf.aviso(f"{r['ref']}, titulo: tiene {len(plano(titulo))} caracteres; en celular se verá largo (ideal < 90).")
        slides.append(s)
    if not slides:
        inf.error("Slides: no hay ningún slide visible. Debe haber al menos uno.")
        slides = [{"ref": "Slides", "titulo": Markup(f"{FALTA_INI}Slides{FALTA_FIN}"), "descripcion": "", "imagen": "",
                   "imagen_alt": "", "cta_texto": "", "cta_enlace": "", "cta2_texto": "", "cta2_enlace": ""}]

    # ---- Productos
    productos = []
    for r in productos_r:
        nombre = campo(r, "nombre", True)
        p = {"ref": r["ref"], "marca": campo(r, "marca", True), "nombre": nombre, "nombre_plano": plano(nombre),
             "imagen": imagen(r, "imagen", True), "imagen_alt": r.get("imagen_alt") or plano(nombre),
             "enlace": enlace(r, "enlace"), "etiqueta": campo(r, "etiqueta", False)}
        p["beneficio"] = rico(campo(r, "beneficio", False)) or t["producto_beneficio_defecto"]
        if len(nombre) > 90:
            inf.aviso(f"{r['ref']}, nombre: muy largo ({len(nombre)} caracteres).")
        productos.append(p)
    if not productos:
        inf.error("Productos: no hay ningún producto visible.")

    beneficios = [{"ref": r["ref"], "icono": campo(r, "icono", True), "titulo": rico(campo(r, "titulo", True)),
                   "texto": rico(campo(r, "texto", False))} for r in beneficios_r]
    for b in beneficios:
        if len(b["icono"]) > 5:
            inf.aviso(f"{b['ref']}, icono: «{b['icono']}» es largo; el círculo admite ~4 caracteres.")
    pasos = [{"ref": r["ref"], "titulo": rico(campo(r, "titulo", True)), "texto": rico(campo(r, "texto", False))}
             for r in pasos_r]
    faq = [{"ref": r["ref"], "pregunta": rico(campo(r, "pregunta", True)), "respuesta": rico(campo(r, "respuesta", True))}
           for r in faq_r]
    for nombre_sec, lista in (("Beneficios", beneficios), ("Pasos", pasos), ("FAQ", faq)):
        if not lista:
            inf.aviso(f"{nombre_sec}: no hay filas visibles; la sección no se muestra.")

    # ---- Legales: texto íntegro, sin variables ni formato
    legales = []
    for r in legales_r:
        if not r.get("texto"):
            inf.aviso(f"{r['ref']}: fila de legales sin texto, se omite.")
            continue
        legales.append({"ref": r["ref"], "texto": literal(r["texto"])})
    if not legales:
        legales = [{"ref": "Legales", "texto": Markup(f"{FALTA_INI}Legales, texto{FALTA_FIN}")}]
        inf.error("Legales: no hay texto. Pega los términos y condiciones completos aprobados.")
    else:
        todo = " ".join(r["texto"] for r in legales_r if r.get("texto")).lower()
        for nombre_dato, valor in (("código", c["codigo"]), ("porcentaje", c["porcentaje"]), ("tope", c["tope"]), ("mínimo", c["minimo"])):
            if valor and FALTA_INI not in valor and valor.lower() not in todo:
                inf.aviso(f"Legales: no mencionan el {nombre_dato} «{valor}» de Config. Revisa que sean los legales de esta campaña.")

    ctx = {"c": c, "t": t, "slides": slides, "datos_hero": datos_hero, "productos": productos,
           "beneficios": beneficios, "pasos": pasos, "faq": faq, "legales": legales}
    return ctx, inf, urls


# --------------------------------------------------------------------------- verificación en red

def verificar_urls(urls: dict[str, tuple[bool, str]], inf: Informe) -> None:
    def una(item):
        url, (es_img, donde) = item
        cab = {"User-Agent": "Mozilla/5.0 (generador landing BCP Juntoz)"}
        for metodo in ("HEAD", "GET"):
            try:
                req = urllib.request.Request(url, method=metodo, headers=cab)
                with urllib.request.urlopen(req, timeout=15) as resp:
                    tipo = resp.headers.get("Content-Type", "")
                    if es_img and not tipo.startswith("image/"):
                        return f"{donde}: «{url}» no devuelve una imagen (Content-Type: {tipo or 'vacío'})."
                    return None
            except urllib.error.HTTPError as e:
                if metodo == "HEAD" and e.code in (403, 405, 501):
                    continue
                return f"{donde}: «{url}» responde {e.code}."
            except Exception as e:  # red, DNS, TLS
                return f"{donde}: no se pudo abrir «{url}» ({type(e).__name__}: {e})."
        return None

    with ThreadPoolExecutor(max_workers=8) as ex:
        for res in ex.map(una, urls.items()):
            if res:
                inf.error(res)


# --------------------------------------------------------------------------- render

def renderizar(ctx: dict, modo: str, comentario: str) -> str:
    env = Environment(loader=FileSystemLoader(str(BASE)), autoescape=True, undefined=StrictUndefined,
                      trim_blocks=False, lstrip_blocks=False, keep_trailing_newline=True)
    env.globals["df"] = (lambda ref: Markup(f' data-field="{escape(ref)}"')) if modo == "preview" else (lambda ref: "")
    html = env.get_template(PLANTILLA).render(meta={"comentario": comentario}, **ctx)
    html = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", html).lstrip()
    return html


def revisar_final(html: str, inf: Informe) -> None:
    """Controles sobre el HTML final: compatible con Magento y sin nada peligroso."""
    etiquetas = " ".join(re.findall(r"<[^>]+>", html))
    if re.search(r"<script\b", html, re.I) or re.search(r"\son[a-z]+\s*=", etiquetas, re.I) or "javascript:" in etiquetas.lower():
        inf.error("HTML final: contiene JavaScript. No se publica.")
    if "data-field" in html:
        inf.error("HTML final: quedaron atributos data-field.")
    if FALTA_INI in html:
        inf.error("HTML final: quedan campos obligatorios sin completar.")
    if RE_SECRETO.search(html):
        inf.error("HTML final: contiene algo que parece una credencial. No se publica.")
    cuerpo = re.sub(r"<style>.*?</style>", "", html, flags=re.S)
    if RE_TOKEN.search(cuerpo):
        inf.error(f"HTML final: quedó una variable sin reemplazar ({RE_TOKEN.search(cuerpo).group(0)}).")
    if len(re.findall(r"<h1\b", html)) != 1:
        inf.error("HTML final: debe tener exactamente un H1.")
    estilo = re.search(r"<style>(.*?)</style>", html, re.S).group(1)
    sin_comentarios = re.sub(r"/\*.*?\*/", "", estilo, flags=re.S)
    for selector in re.findall(r"([^{}]+)\{", sin_comentarios):
        selector = selector.strip()
        if selector.startswith("@"):
            continue
        for parte in selector.split(","):
            if not parte.strip().startswith("#bcp-landing"):
                inf.error(f"HTML final: el selector CSS «{parte.strip()}» no está encapsulado en #bcp-landing.")


def marcar_faltantes(html: str) -> str:
    """En la vista previa resalta los ⟦FALTA⟧ que están en texto (no dentro de atributos)."""
    patron = re.compile(r"(<[^>]*>)|" + re.escape(FALTA_INI) + r"([^" + FALTA_FIN + r"]*)" + re.escape(FALTA_FIN))
    return patron.sub(lambda m: m.group(1) or f'<mark class="pv-falta">FALTA: {m.group(2)}</mark>', html)


def documento_preview(cuerpo: str, inf: Informe, origen: str, publicable: bool) -> str:
    def lista(items, clase):
        return "".join(f'<li class="{clase}">{escape(i)}</li>' for i in items)
    estado = ("LISTA PARA PUBLICAR" if publicable else "NO SE PUEDE PUBLICAR: corrige los errores")
    return f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Vista previa · Landing BCP</title>
<style>
body{{margin:0;background:#E9EBEF;font-family:system-ui,Arial,sans-serif}}
.pv-panel{{background:#111827;color:#F9FAFB;padding:16px;font-size:13px;line-height:1.5}}
.pv-panel h1{{font-size:15px;margin:0 0 6px}} .pv-panel p{{margin:0 0 8px;color:#D1D5DB}}
.pv-ok{{color:#86EFAC}} .pv-ko{{color:#FCA5A5}} .pv-panel ul{{margin:6px 0 0;padding-left:18px}}
.pv-err{{color:#FCA5A5}} .pv-avi{{color:#FDE68A}}
.pv-panel details summary{{cursor:pointer}}
.pv-tienda{{background:#fff;border-bottom:1px solid #D1D5DB;padding:14px 16px;color:#6B7280;font-size:12px;text-align:center}}
.pv-falta{{background:#FDE047;color:#7F1D1D;font-weight:700;padding:0 4px}}
[data-field]:hover{{outline:2px dashed #F97316;outline-offset:2px;position:relative}}
[data-field]:hover::after{{content:attr(data-field);position:absolute;left:0;top:100%;z-index:9;background:#F97316;color:#fff;font:600 11px/1.4 system-ui;padding:2px 6px;border-radius:4px;white-space:nowrap;letter-spacing:0;text-transform:none}}
</style></head><body>
<div class="pv-panel">
  <h1>Vista previa · Landing BCP · <span class="{'pv-ok' if publicable else 'pv-ko'}">{estado}</span></h1>
  <p>Origen: {escape(origen)} · {len(inf.errores)} errores · {len(inf.avisos)} avisos. Pasa el mouse sobre un texto para ver de qué celda del Sheet sale. Este panel y las barras grises no se publican.</p>
  {'<details open><summary>Errores</summary><ul>' + lista(inf.errores, 'pv-err') + '</ul></details>' if inf.errores else ''}
  {'<details><summary>Avisos</summary><ul>' + lista(inf.avisos, 'pv-avi') + '</ul></details>' if inf.avisos else ''}
</div>
<div class="pv-tienda">Header de Juntoz (no forma parte de la landing)</div>
{marcar_faltantes(cuerpo)}
<div class="pv-tienda">Footer de Juntoz (no forma parte de la landing)</div>
</body></html>
"""


# --------------------------------------------------------------------------- plantilla xlsx

def crear_plantilla(ruta: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    azul = PatternFill("solid", fgColor="0030B2")
    blanco = Font(color="FFFFFF", bold=True)
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    lineas = [
        "Plantilla de la landing BCP en Juntoz",
        "",
        "1. Completa Config (datos de la campaña) y Legales (texto completo aprobado, sin resumir).",
        "2. Edita Textos, Slides, Productos, Beneficios, Pasos y FAQ. visible = SI / NO; orden = 1, 2, 3...",
        f"3. En los textos puedes usar variables: {TOKENS_AYUDA}. Se reemplazan con los datos de Config.",
        "4. **texto** pone el texto en negrita (en el título del slide lo pinta de azul).",
        "5. No renombres pestañas ni encabezados. Las celdas vacías opcionales ocultan su elemento.",
        "6. Legales: una fila por párrafo. Se publican tal cual (sin variables ni formato).",
        "7. Para publicar: legales_aprobados = SI, con quién aprobó y la fecha.",
        "8. Imágenes: URLs https de los dominios permitidos en Config. Slides en formato cuadrado (1:1).",
        "9. Luego corre el generador (ver README): primero vista previa, después publicar.",
    ]
    for i, l in enumerate(lineas, start=1):
        ins.cell(row=i, column=1, value=l)
    ins["A1"].font = Font(bold=True, size=14, color="0030B2")
    ins.column_dimensions["A"].width = 120

    def hoja(nombre, filas, anchos):
        ws = wb.create_sheet(nombre)
        ws.append(COLUMNAS[nombre])
        for c in ws[1]:
            c.fill, c.font = azul, blanco
        for f in filas:
            ws.append(f)
        for i, w in enumerate(anchos):
            ws.column_dimensions[chr(65 + i)].width = w
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        if "visible" in COLUMNAS[nombre]:
            col = chr(65 + COLUMNAS[nombre].index("visible"))
            dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}200")
        return ws

    cfg = hoja("Config", [[k, d, a + (" (obligatorio)" if o else " (opcional)")] for k, o, _, d, a in CONFIG], [28, 60, 90])
    fila_leg = [k for k, *_ in CONFIG].index("legales_aprobados") + 2
    dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=False)
    cfg.add_data_validation(dv)
    dv.add(f"B{fila_leg}")
    hoja("Textos", [[k, d, a + (" (obligatorio)" if o else " (opcional)")] for k, o, d, a in TEXTOS], [30, 80, 70])
    hoja("Slides", EJEMPLO["Slides"], [8, 9, 50, 60, 50, 30, 22, 30, 22, 30])
    hoja("Productos", EJEMPLO["Productos"], [8, 9, 14, 45, 60, 30, 70, 25, 14])
    hoja("Beneficios", EJEMPLO["Beneficios"], [8, 9, 8, 30, 50])
    hoja("Pasos", EJEMPLO["Pasos"], [8, 9, 30, 60])
    hoja("FAQ", EJEMPLO["FAQ"], [8, 9, 45, 80])
    hoja("Legales", EJEMPLO["Legales"], [8, 140])
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


# --------------------------------------------------------------------------- comandos

def obtener_datos(args) -> tuple[dict, str]:
    if args.fuente == "xlsx":
        if not args.archivo:
            sys.exit("Indica el archivo: --archivo ruta/al/archivo.xlsx")
        return leer_xlsx(Path(args.archivo)), f"archivo {Path(args.archivo).name}"
    sheet_id = args.sheet_id or os.environ.get("BCP_SHEET_ID")
    cred = args.credenciales or os.environ.get("GOOGLE_APPLICATION_CREDENTIALS")
    if not sheet_id or not cred:
        sys.exit("Para leer Google Sheets define BCP_SHEET_ID y GOOGLE_APPLICATION_CREDENTIALS "
                 "(o usa --sheet-id y --credenciales). Alternativa sin credenciales: --fuente xlsx --archivo descarga.xlsx")
    return leer_sheets(sheet_id, cred), "Google Sheets"


def imprimir_informe(inf: Informe) -> None:
    for e in inf.errores:
        print(f"  ERROR  {e}")
    for a in inf.avisos:
        print(f"  aviso  {a}")
    print(f"  {len(inf.errores)} errores, {len(inf.avisos)} avisos.")


def generar(args, publicar: bool) -> int:
    datos, origen = obtener_datos(args)
    hoy = dt.date.fromisoformat(args.hoy) if args.hoy else None
    ctx, inf, urls = construir(datos, hoy)
    if args.verificar_urls:
        print(f"Verificando {len(urls)} URLs...")
        verificar_urls(urls, inf)
    elif publicar:
        inf.aviso("No se verificó que enlaces e imágenes respondan (usa --verificar-urls antes de publicar).")
    ahora = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    comentario = (f"Landing BCP · campaña {ctx['c']['campana']} · generada {ahora} desde {origen}. "
                  "No editar a mano: cambiar el Sheet y volver a generar.")
    final = renderizar(ctx, "final", comentario)
    revisar_final(final, inf)
    preview = documento_preview(renderizar(ctx, "preview", comentario), inf, origen, not inf.errores)

    salida = Path(args.salida) if args.salida else DIR_SALIDA
    salida.mkdir(parents=True, exist_ok=True)
    (salida / "bcp-landing.preview.html").write_text(preview, encoding="utf-8")
    print(f"Vista previa: {salida / 'bcp-landing.preview.html'}")
    imprimir_informe(inf)

    if not publicar:
        return 1 if inf.errores else 0
    if inf.errores:
        print("NO SE PUBLICÓ: corrige los errores en el Sheet y vuelve a correr.")
        return 1
    destino_dir = Path(args.publicado) if args.publicado else DIR_PUBLICADO
    respaldos = Path(args.respaldos) if args.respaldos else DIR_RESPALDOS
    destino_dir.mkdir(parents=True, exist_ok=True)
    destino = destino_dir / NOMBRE_FINAL
    if destino.exists():
        anterior = destino.read_text(encoding="utf-8")
        respaldos.mkdir(parents=True, exist_ok=True)
        sello = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.copy2(destino, respaldos / f"bcp-landing_{sello}.html")
        diff = difflib.unified_diff(anterior.splitlines(), final.splitlines(), "anterior", "nueva", lineterm="")
        (respaldos / f"bcp-landing_{sello}.diff").write_text("\n".join(diff), encoding="utf-8")
        print(f"Respaldo de la versión anterior: {respaldos / f'bcp-landing_{sello}.html'}")
    destino.write_text(final, encoding="utf-8")
    print(f"PUBLICABLE: {destino} ({len(final.encode()) // 1024} KB). Pégalo en el bloque/página CMS de Magento.")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Generador de la landing BCP de Juntoz desde Google Sheets.")
    sub = ap.add_subparsers(dest="comando", required=True)
    p_pl = sub.add_parser("plantilla", help="Crea la plantilla .xlsx para importar al Google Sheet.")
    p_pl.add_argument("--archivo", default=str(DIR_SALIDA / "BCP_Google_Sheets_Plantilla.xlsx"))
    for nombre, ayuda in (("previsualizar", "Genera la vista previa local con el informe de validación."),
                          ("publicar", "Valida y genera publicado/bcp-landing.html, con respaldo de la versión anterior.")):
        p = sub.add_parser(nombre, help=ayuda)
        p.add_argument("--fuente", choices=["sheets", "xlsx"], default="sheets")
        p.add_argument("--archivo", help="Ruta del .xlsx (con --fuente xlsx).")
        p.add_argument("--sheet-id", help="ID del Google Sheet (por defecto: variable BCP_SHEET_ID).")
        p.add_argument("--credenciales", help="JSON de la cuenta de servicio (por defecto: GOOGLE_APPLICATION_CREDENTIALS).")
        p.add_argument("--verificar-urls", action="store_true", help="Comprueba que enlaces e imágenes respondan.")
        p.add_argument("--hoy", help=argparse.SUPPRESS)
        p.add_argument("--salida", help=argparse.SUPPRESS)
        p.add_argument("--publicado", help=argparse.SUPPRESS)
        p.add_argument("--respaldos", help=argparse.SUPPRESS)
    args = ap.parse_args(argv)
    if args.comando == "plantilla":
        crear_plantilla(Path(args.archivo))
        print(f"Plantilla creada: {args.archivo}")
        return 0
    return generar(args, publicar=args.comando == "publicar")


if __name__ == "__main__":
    sys.exit(main())
